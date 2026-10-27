#!/usr/bin/env python3
"""
readers.py — Format-specific document readers for the CLIQ parser.

This module is the *only* place that knows how to turn a file on disk into
text. The extraction logic (doc_parser.py) never imports a PDF/DOCX library —
it only consumes the page list these readers return. New formats are added by
registering another reader function; the core extractor is untouched.

A reader takes a file path and returns ``list[str]`` — one string per logical
page. Plain text and DOCX have no real pagination, so they return a single
page; PDFs return one entry per page (useful for running-header detection).

CONSTRAINTS:
  - Fully local. No network access, ever.
  - Optional heavy dependencies (pypdf, python-docx) are imported lazily so the
    rest of CLIQ works without them. A missing dependency raises a clear,
    actionable ReaderDependencyError rather than a bare ImportError.
"""

from __future__ import annotations

import os
from typing import Callable

# A reader: path -> list of page texts.
Reader = Callable[[str], list]


class ReaderError(Exception):
    """Base class for all reader failures (corrupt file, bad format, etc.)."""


class UnsupportedFormatError(ReaderError):
    """Raised when no reader is registered for a file extension."""


class ReaderDependencyError(ReaderError):
    """Raised when the optional library for a format is not installed."""


# ---------------------------------------------------------------------------
# Concrete readers
# ---------------------------------------------------------------------------
def read_txt(path: str) -> list[str]:
    """Read a plain-text file. Splits on form-feed (\\f) if the file uses it
    as a page separator; otherwise the whole file is one page."""
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        text = fh.read()
    return text.split("\f") if "\f" in text else [text]


def read_pdf(path: str) -> list[str]:
    """Read a PDF, returning one text string per page. Requires `pypdf`
    (falls back to the legacy `PyPDF2` name if that is what is installed)."""
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError as exc:
            raise ReaderDependencyError(
                "Reading PDF files requires 'pypdf'. Install it with: "
                "pip install pypdf"
            ) from exc
    try:
        reader = PdfReader(path)
        return [(page.extract_text() or "") for page in reader.pages]
    except ReaderError:
        raise
    except Exception as exc:  # corrupt/encrypted PDF, etc.
        raise ReaderError(f"Failed to read PDF '{path}': {exc}") from exc


def read_docx(path: str) -> list[str]:
    """Read a DOCX, flattening paragraphs and table cells into one page.
    Requires `python-docx`."""
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise ReaderDependencyError(
            "Reading DOCX files requires 'python-docx'. Install it with: "
            "pip install python-docx"
        ) from exc
    try:
        document = docx.Document(path)
    except Exception as exc:
        raise ReaderError(f"Failed to read DOCX '{path}': {exc}") from exc

    parts: list[str] = [p.text for p in document.paragraphs]
    # Table rows are common in command references — keep cell order, tab-joined
    # so the extractor can still see "command<TAB>description" pairs.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return ["\n".join(parts)]


# ---------------------------------------------------------------------------
# Registry — extensible mapping of file extension -> reader
# ---------------------------------------------------------------------------
_READERS: dict[str, Reader] = {
    ".txt": read_txt,
    ".text": read_txt,
    ".pdf": read_pdf,
    ".docx": read_docx,
}


def register_reader(extension: str, reader: Reader) -> None:
    """Register (or override) the reader for a file extension. This is how new
    formats are plugged in without modifying the core extractor."""
    _READERS[extension.lower()] = reader


def supported_extensions() -> list[str]:
    return sorted(_READERS.keys())


def read_document(path: str) -> list[str]:
    """Dispatch to the reader for `path`'s extension and return its pages.

    Raises UnsupportedFormatError, ReaderDependencyError, or ReaderError.
    """
    if not os.path.isfile(path):
        raise ReaderError(f"File not found: {path}")
    ext = os.path.splitext(path)[1].lower()
    reader = _READERS.get(ext)
    if reader is None:
        raise UnsupportedFormatError(
            f"No reader for '{ext or '(no extension)'}'. "
            f"Supported: {', '.join(supported_extensions())}"
        )
    return reader(path)
